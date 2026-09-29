"""Daily report: everything FMCSA's "Daily Registration Decisions" page used to publish, as one Excel file.

That page stopped listing documents on 2026-05-20 and its PDFs are no longer served; Motus open data carries
the same decisions. Each row is enriched with the Motus carrier record (name, phone, address) and the census
(email, trucks, officer). Sent to Telegram once per day (collector_runs rows with source 'daily_report').
"""
import logging
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.collectors.fmcsa import (CENSUS, CHUNK, MOTUS_AUTHORITY, MOTUS_AUTHORITY_DAILY, MOTUS_CARRIER,
                                  MOTUS_CARRIER_DAILY, MOTUS_INSURANCE_HISTORY, MOTUS_ORDERS, MOTUS_ORDERS_DAILY,
                                  FmcsaCollector, _in, docket_id)
from app.config import get_settings
from app.models import CollectorRun

log = logging.getLogger(__name__)
SOURCE = "daily_report"
OUT_DIR = Path("reports")
# FMCSA works on US Eastern time; by mid-morning there the previous day's daily-difference sets are published.
FMCSA_TZ, READY_HOUR = ZoneInfo("America/New_York"), 10
ADMIN_FIX = "Admin correction (MOTUS Issue #220)"
COLS = ["Date", "Event", "Authority type", "Docket", "USDOT", "Legal name", "DBA", "Phone", "Email", "Officer",
        "Street", "City", "State", "Zip", "Trucks", "Drivers", "Authority status", "BIPD on file ($)", "Insurer",
        "Suspension effective"]
NOTES = {
    "4 Regular Routes (passenger)": "Motus doesn't mark regular-route authority; this is every passenger-carrier grant.",
    "5 Name Changes": "Motus publishes no name-change decisions; these are insurance filings re-issued under a new name.",
    "6 Reinstatements": "'Admin correction (MOTUS Issue #220)' rows are FMCSA fixing its own data, not real reinstatements.",
    "7 Transfers": "Motus open data publishes no authority transfers.",
    "8 Broker-FF Financial Security": "Broker/freight-forwarder suspension notices, plus surety bond / trust fund "
                                      "cancellations taking effect up to 30 days after the period.",
}


def categorize(auth: list[dict], orders: list[dict], ins: list[dict], lo: str, hi: str) -> dict[str, list[dict]]:
    """Split Motus rows into the old page's 8 categories. lo/hi: YYYYMMDD. Rows get 'event' and a 'date'."""
    reason = lambda r: r.get("reason") or ""  # noqa: E731
    typ = lambda r: r.get("op_auth_type") or ""  # noqa: E731
    ev = lambda r, event, d: {**r, "event": event, "date": d}  # noqa: E731
    granted = [ev(r, "Granted", r.get("status_change_date")) for r in auth if reason(r).lower() == "granted"]
    return {
        "1 Daily Register": [ev(r, reason(r), r.get("status_change_date")) for r in auth
                             if reason(r) in ("Published to FMCSA Register", "Revoked", "Out of Service")
                             or reason(r).startswith("Involuntary Suspension")]
                            + [ev(r, r.get("order1_type_desc"), r.get("order1_serve_date")) for r in orders],
        "2 Certificates of Authority": [r for r in granted if not typ(r).startswith("Mexico") and "Passengers" not in typ(r)],
        "3 OP-2 MX Commercial Zone": [r for r in granted if typ(r).startswith("Mexico")],
        "4 Regular Routes (passenger)": [r for r in granted if "Passengers" in typ(r)],
        "5 Name Changes": [ev(r, "Insurance filing re-issued under new name (NAMECHG)", r.get("cancl_effective_date"))
                           for r in ins if r.get("filing_status_reason") == "NAMECHG"
                           and lo <= (r.get("cancl_effective_date") or "") <= hi],
        "6 Reinstatements": [ev(r, ADMIN_FIX if "#220" in reason(r) else reason(r),
                                r.get("status_change_date")) for r in auth if reason(r).startswith("Reinstated")],
        "7 Transfers": [],
        "8 Broker-FF Financial Security": [
            ev(r, r.get("order1_type_desc"), r.get("order1_serve_date")) for r in orders
            if typ(r).startswith(("Broker", "Freight Forwarder")) and "Involuntary" in (r.get("order1_type_desc") or "")]
            + [{**ev(r, f"{r.get('ins_type_desc')} cancellation effective {r.get('cancl_effective_date')}",
                     r.get("cancl_effective_date")), "op_auth_type": "(bond / trust fund)"}
               for r in ins if r.get("ins_type_desc") in ("SURETY", "TRUST FUND") and r.get("filing_status_reason") == "CANCEL"],
    }


def build(lo: date, hi: date, path: Path, collector: FmcsaCollector | None = None) -> dict[str, int]:
    """Fetch the decisions for lo..hi, write the workbook, return rows per category."""
    get = (collector or FmcsaCollector())._get
    L, H = f"{lo:%Y%m%d}", f"{hi:%Y%m%d}"

    def both(ds, daily, where, keys):  # history + daily-difference set (newest day), one row per decision
        return list({tuple(docket_id(r.get(k)) for k in keys): r for r in get(ds, where) + get(daily, where)}.values())

    auth = both(MOTUS_AUTHORITY, MOTUS_AUTHORITY_DAILY, f"status_change_date between '{L}' and '{H}'",
                ("usdot_number", "docket_number", "op_auth_type", "reason", "status_change_date"))
    orders = both(MOTUS_ORDERS, MOTUS_ORDERS_DAILY, f"order1_serve_date between '{L}' and '{H}'",
                  ("usdot_number", "docket_number", "order1_type_desc", "order1_serve_date"))
    ins = get(MOTUS_INSURANCE_HISTORY, f"cancl_effective_date between '{L}' and '{hi + timedelta(days=30):%Y%m%d}'")
    sheets = categorize(auth, orders, ins, L, H)

    dockets = sorted({docket_id(r.get("docket_number")) for rows in sheets.values() for r in rows} - {""})
    dots = sorted({r.get("usdot_number") for rows in sheets.values() for r in rows if (r.get("usdot_number") or "").isdigit()})
    carrier, census = {}, {}
    for i in range(0, len(dockets), CHUNK):
        for ds in (MOTUS_CARRIER, MOTUS_CARRIER_DAILY):  # daily last, so it wins
            carrier |= {r["docket_number"]: r for r in get(ds, f"docket_number in ({_in(dockets[i:i + CHUNK])})")}
    for i in range(0, len(dots), CHUNK):
        census |= {r["dot_number"]: r for r in get(CENSUS, f"dot_number in ({_in(dots[i:i + CHUNK])})")}

    def row(r):
        k, dot = docket_id(r.get("docket_number")), r.get("usdot_number") or ""
        m, s, d = carrier.get(k, {}), census.get(dot, {}), r.get("date") or ""
        bipd = m.get("bipd_file")
        return [f"{d[:4]}-{d[4:6]}-{d[6:8]}" if len(d) >= 8 else d, r.get("event"), r.get("op_auth_type") or m.get("op_auth_type"),
                k, dot, m.get("legal_name") or s.get("legal_name"), s.get("dba_name"), m.get("bus_telno") or s.get("phone"),
                (s.get("email_address") or "").lower() or None, s.get("company_officer_1"),
                m.get("bus_street_po") or s.get("phy_street"), m.get("bus_city") or s.get("phy_city"),
                m.get("bus_state_code") or s.get("phy_state"), m.get("bus_zip_code") or s.get("phy_zip"),
                s.get("power_units"), s.get("total_drivers"), m.get("op_auth_status"),
                float(bipd) if bipd else None, r.get("insurance_company_name"), r.get("order1_effective_date")]

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append([f"FMCSA daily decisions {lo} to {hi}, from Motus open data (data.transportation.gov)"])
    summary.append(["Category", "Records", "Note"])
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(COLS)
        for r in sorted(rows, key=lambda r: r.get("date") or "", reverse=True):
            ws.append(row(r))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
        summary.append([name, len(rows), NOTES.get(name, "")])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    real = {name: [r for r in rows if r.get("event") != ADMIN_FIX] for name, rows in sheets.items()}
    counts = {name: len(rows) for name, rows in real.items()}
    counts["companies"] = len({r.get("usdot_number") for rows in real.values() for r in rows} - {None, ""})
    return counts


def send_daily_report(db: Session, day: date | None = None, once: bool = False, notifier=None) -> dict | None:
    """Build the report for `day` (default yesterday) and send it to Telegram. once: skip a day already sent."""
    from app.notify import TelegramNotifier

    s = get_settings()
    now = datetime.now(FMCSA_TZ)
    if day is None:
        day = now.date() - timedelta(days=1)
        if once and now.hour < READY_HOUR:
            log.info("daily report for %s: FMCSA data not complete before %d:00 US Eastern", day, READY_HOUR)
            return None
    if once and db.scalar(select(CollectorRun.id).where(
            CollectorRun.source == SOURCE, CollectorRun.status == "success", CollectorRun.cursor == day)):
        log.info("daily report for %s already sent", day)
        return None
    run = CollectorRun(source=SOURCE, cursor=day, stats={})
    db.add(run)
    db.commit()
    try:
        path = OUT_DIR / f"FMCSA-decisions-{day}.xlsx"
        counts = build(day, day, path)
        if notifier is None and s.telegram_bot_token and s.telegram_chat_id:
            notifier = TelegramNotifier(s.telegram_bot_token, s.telegram_chat_id)
        if notifier:
            caption = "\n".join([f"📋 <b>FMCSA daily decisions — {day:%a %d %b %Y}</b>",
                                  f"{counts['companies']} companies", ""]
                                 + [f"• {name[2:]}: {n}" for name, n in counts.items() if n and name != "companies"])
            notifier.send_document(path, caption)
        # not 'success' unless sent, so `--once` retries after Telegram is configured
        run.status = "success" if notifier else "built"
        run.stats = {"counts": counts, "file": str(path), "sent": bool(notifier)}
    except Exception as e:
        run.status, run.error = "failed", f"{type(e).__name__}: {e}"[:2000]
        log.exception("daily report for %s failed", day)
    finally:
        run.finished_at = datetime.now(timezone.utc)
        db.commit()
    log.info("daily report %s: %s %s", day, run.status, run.stats)
    return run.stats
